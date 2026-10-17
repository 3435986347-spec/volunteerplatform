#!/usr/bin/env python3
"""把二进制的需求原件转成 agent 能直接读的 markdown，写到 `知识库/需求/`。

    需求表.md   《小程序设想【第十版】.xlsx》的「前端」工作表，一行一节，标题就是表格行号
               （代码注释、迁移、知识库里写的「Row 47」就是这里的「Row 47」）
    原型.md     《志愿平台设想【第十版】.pdf》127 页，一页一节（「原型 P83」就是这里的「P83」）

原件留在 `文档/`，这里只是文字版：单元格里的图片、原型的版式与配图都没有转过来，要看图仍得打开原件。
**手机号、身份证号一律打码**——原型里有协会联系人的真实号码，知识库会推到单独的仓库里。

依赖：openpyxl、pypdf（`pip3 install openpyxl pypdf`）。原件很少变，变了重跑一次即可：

    python3 tools/convert_requirements.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

XLSX = "文档/小程序设想【第十版】.xlsx"
PDF = "文档/志愿平台设想【第十版】.pdf"
OUT = "知识库/需求"

# 「前端」工作表的列（第 1 行表头；后台段在第 58 行另有一行表头，列含义相同）
COLUMNS = [(3, "前端信息"), (4, "前端备注"), (6, "后端功能要求备注"), (7, "后端功能备注")]
SECTION_COL, FEATURE_COL = 1, 2

_PHONE = re.compile(r"(?<!\d)(1[3-9]\d)\d{4}(\d{4})(?!\d)")
_ID_CARD = re.compile(r"(?<![\dA-Za-z])(\d{6})\d{8}(\d{3}[\dXx])(?![\dA-Za-z])")

HEADER = "<!-- 由 tools/convert_requirements.py 从原件转出，不要手改；原件变了重跑脚本。 -->\n\n"


def mask(text: str) -> str:
    text = _ID_CARD.sub(lambda m: m.group(1) + "********" + m.group(2), text)
    return _PHONE.sub(lambda m: m.group(1) + "****" + m.group(2), text)


def _cell(ws, row: int, col: int) -> str:
    v = ws.cell(row, col).value
    return "" if v is None else mask(str(v).strip())


def _hard_breaks(text: str) -> str:
    return "  \n".join(line.rstrip() for line in text.splitlines())


def convert_xlsx(root: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(root / XLSX, data_only=True)
    ws = wb["前端"]
    out = [HEADER, "# 需求表（《小程序设想【第十版】.xlsx》「前端」工作表）\n\n",
           f"原件：`{XLSX}`。**标题里的行号就是表格行号**——代码注释、迁移、协会答复里写的「Row 47」即本页的「Row 47」。\n"
           "第 2–49 行是小程序（首页 / 社区 / 我的），第 58–79 行是「后台相关要求」，第 82–86 行是「监管后台」（未建）。\n"
           "单元格里的图片没有转过来；`【】` 是原件里的勾选框记号。手机号已打码。\n\n"
           "各行在系统里落在哪个领域、做到什么程度，看 `知识库/领域/*.md` 的「需求出处」一节。\n\n"]
    index, body = [], []
    section = ""
    for r in range(2, ws.max_row + 1):
        sec = _cell(ws, r, SECTION_COL)
        feature = _cell(ws, r, FEATURE_COL)
        fields = [(label, _cell(ws, r, c)) for c, label in COLUMNS]
        if sec and sec != section and sec not in ("功能",):
            section = sec
            body.append(f"\n## {section}\n")
        if feature == "功能" or not (feature or any(v for _, v in fields)):
            continue                     # 后台段的第二行表头、空行
        title = feature.splitlines()[0][:40] if feature else "（无标题）"
        index.append(f"| {r} | {section} | {title.replace('|', '/')} |\n")
        body.append(f"\n### Row {r} · {title}\n")
        if feature and (len(feature) > 40 or "\n" in feature):
            body.append(f"\n**功能**  \n{_hard_breaks(feature)}\n")
        for label, value in fields:
            if value:
                body.append(f"\n**{label}**  \n{_hard_breaks(value)}\n")
    out.append("| Row | 板块 | 功能 |\n|---|---|---|\n")
    out.extend(index)
    out.extend(body)
    return "".join(out)


def convert_pdf(root: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(root / PDF))
    out = [HEADER, "# 原型（《志愿平台设想【第十版】.pdf》）\n\n",
           f"原件：`{PDF}`，共 {len(reader.pages)} 页。**标题里的页码就是 PDF 页码**——材料里写的「原型 P83」即本页的「P83」。\n"
           "这是从版式里抽出来的文字：同一屏的多个状态会挤在一起、左右两栏会交错，配图与箭头都没有；"
           "要确认布局或图里的标注，打开原件对应页。手机号已打码。\n"]
    for i, page in enumerate(reader.pages, 1):
        text = mask((page.extract_text() or "").strip())
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        out.append(f"\n## P{i}\n\n```text\n{text or '（这一页抽不出文字，多半是整页图片）'}\n```\n")
    return "".join(out)


def main() -> int:
    root = Path(".").resolve()
    if not (root / XLSX).exists():
        print(f"找不到 {XLSX}——请在文档仓库根目录运行", file=sys.stderr)
        return 2
    out_dir = root / OUT
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "需求表.md").write_text(convert_xlsx(root), encoding="utf-8")
    (out_dir / "原型.md").write_text(convert_pdf(root), encoding="utf-8")
    print(f"写了 {OUT}/需求表.md、{OUT}/原型.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
